"""
Build the portal data from the VCT Sector Operating Model workbook.

Usage:
    python scripts/build_data.py [path/to/workbook.xlsx]

Writes portal/data/vct-data.js (window.VCT_DATA) so the portal also works when
index.html is opened directly. The live server (scripts/live_server.py) imports
build() from here and serves the same data at /api/data.

Everything the portal shows is calculated from the raw input columns, not from
the workbook's formula results. The workbook was built in Google Sheets and
several formulas (allocation %, project totals, project leads) only recalculate
there; desktop Excel keeps their old values. Reading the inputs keeps the
portal correct whichever program saves the file.
"""
import datetime as dt
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CONFIG = Path(os.environ.get("VCT_PORTAL_CONFIG") or ROOT / "portal.config.json")
OUT = ROOT / "portal" / "data" / "vct-data.js"

NOT_SET = "Not set"
LEAD_SUFFIX = re.compile(r"\s+—\s+(Sector|Project) Lead$")
GOOGLE_ONLY = "__xludf.DUMMYFUNCTION"
HEADCOUNT_SHEET = "Headcount Data"
SUMMARY_SHEET = "Headcount Summary"


def default_workbook():
    if CONFIG.exists():
        cfg = json.loads(CONFIG.read_text(encoding="utf-8"))
        if cfg.get("excelPath"):
            return Path(cfg["excelPath"])
    return ROOT / "Copy of VCT_Sector_Operating_Model.xlsx"


# ---------------------------------------------------------------- cell helpers
def clean(v):
    """Trim strings; treat blanks and the sheet's '—' placeholder as None."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v in ("", "—", "-"):
            return None
    return v


def number(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().rstrip("%")) / (100 if str(v).strip().endswith("%") else 1)
    except ValueError:
        return None


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def billing(v):
    return clean(v) or NOT_SET


def strip_lead(v):
    return LEAD_SUFFIX.sub("", v) if isinstance(v, str) else v


class Grid:
    """A sheet loaded into memory; cell(row, col) is 1-based like Excel."""

    def __init__(self, ws):
        self.rows = [list(r) for r in ws.iter_rows(values_only=True)]

    def cell(self, r, c):
        if r - 1 < len(self.rows):
            row = self.rows[r - 1]
            if c - 1 < len(row):
                return row[c - 1]
        return None

    @property
    def max_row(self):
        return len(self.rows)

    def find_header(self, row, name, start=1):
        row_vals = self.rows[row - 1] if row - 1 < len(self.rows) else []
        for i in range(start - 1, len(row_vals)):
            if clean(row_vals[i]) == name:
                return i + 1
        raise KeyError(f"Column '{name}' not found in header row {row}")


def sheet(wb, name):
    """Sheet by name, ignoring trailing spaces (the workbook has 'Headcount Data ')."""
    for ws in wb.worksheets:
        if ws.title.strip() == name:
            return ws
    raise KeyError(f"Sheet '{name}' not found")


# ---------------------------------------------------------------- readers
def read_operating_model(g):
    out = []
    for r in range(4, g.max_row + 1):
        name = clean(g.cell(r, 1))
        if not name or name == "Dashboard link":
            continue
        out.append({"name": name, "domains": clean(g.cell(r, 2)), "ownership": clean(g.cell(r, 3))})
    return out


def read_sector_heads(g):
    """Overall sheet: sector names in row 5, typed sector heads in row 6."""
    heads = {}
    width = max((len(r) for r in g.rows[:6]), default=0)
    for c in range(2, width + 1):
        sector = clean(g.cell(5, c))
        if sector and sector not in ("Row HC", "Grand total"):
            heads[sector] = clean(g.cell(6, c))
    return heads


def read_projects(g):
    out = []
    for r in range(2, g.max_row + 1):
        sector, name = clean(g.cell(r, 1)), clean(g.cell(r, 2))
        if not sector or not name:
            continue
        aliases = []
        raw = g.cell(r, 6)
        if isinstance(raw, str) and raw.strip().startswith("{"):
            try:
                aliases = [a for a in json.loads(raw).get("aliases", []) if a]
            except ValueError:
                pass
        out.append({
            "sector": sector, "name": name, "lead": clean(g.cell(r, 3)),
            "status": clean(g.cell(r, 4)) or "Active", "aliases": aliases, "row": r,
        })
    return out


def read_employees(gv, gf, projects):
    """gv = values workbook grid, gf = formulas workbook grid (same sheet)."""
    h = lambda name, start=1: gv.find_header(3, name, start)
    col = {k: h(k) for k in [
        "EMP ID", "Employee Name", "Sector", "Project", "First Project %", "Level", "Designation", "Domain",
        "Reports To", "Billed Status", "Earning", "YOE", "Baseline Sector", "Baseline Project", "Baseline Level",
    ]}
    c_email = None
    for name_ in ("Email", "Email ID", "Email Address", "Official Email", "Mail ID", "E-mail"):
        try:
            c_email = h(name_)
            break
        except KeyError:
            pass
    c_alias_from = h("List: Project")
    c_alias_to = h("Grid column (People Mapping)")
    second_start = h("Second Sector")
    col.update({
        "Second Sector": second_start,
        "Second Project": h("Second Project", second_start),
        "Second Billed Status": h("Second Billed Status", second_start),
        "Second Project %": h("Second Project %", second_start),
    })

    # Project name as entered → project column on the people-mapping sheets.
    alias = {}
    for p in projects:
        for a in p["aliases"]:
            alias.setdefault(a, p["name"])
    for r in range(4, gv.max_row + 1):
        a, to = clean(gv.cell(r, c_alias_from)), clean(gv.cell(r, c_alias_to))
        if a and to:
            alias[a] = to

    def v(r, name):
        return gv.cell(r, col[name])

    out = []
    for r in range(4, gv.max_row + 1):
        name = clean(v(r, "Employee Name"))
        if not name:
            continue
        sector = clean(v(r, "Sector"))
        project_raw = clean(v(r, "Project"))
        project = alias.get(project_raw, project_raw)
        level = clean(v(r, "Level"))

        s_sector_raw = clean(v(r, "Second Sector"))
        s_project_raw = clean(v(r, "Second Project"))
        has_second = bool(s_sector_raw and s_project_raw)
        s_sector, s_project = strip_lead(s_sector_raw), strip_lead(s_project_raw)
        # A second entry is a separate assignment only when it differs from the primary one.
        dual = has_second and (s_sector != sector or s_project != project)

        # Allocation %: a typed number wins; otherwise apply the sheet's rule
        # (single project → 100% / 0%, two projects → pending until entered).
        def alloc(name_, single_default):
            raw_f = gf.cell(r, col[name_]) if gf else None
            typed = None if is_formula(raw_f) else number(v(r, name_))
            if typed is not None:
                return typed
            return None if dual else single_default

        first_pct = alloc("First Project %", 1.0)
        base = {k: clean(v(r, "Baseline " + k.title())) for k in ("sector", "project", "level")}
        moved = any(base.values()) and (sector, project_raw, level) != (base["sector"], base["project"], base["level"])

        emp = {
            "id": clean(v(r, "EMP ID")),
            "name": name,
            "sector": sector,
            "project": project,
            "projectRaw": project_raw,
            "firstPct": first_pct,
            "level": level,
            "designation": clean(v(r, "Designation")),
            "domain": clean(v(r, "Domain")) or "Not specified",
            "reportsTo": clean(v(r, "Reports To")),
            "email": clean(gv.cell(r, c_email)) if c_email else None,
            "billing": billing(v(r, "Billed Status")),
            "earning": clean(v(r, "Earning")),
            "yoe": None if clean(v(r, "YOE")) is None else str(clean(v(r, "YOE"))),
            "baseline": base,
            "moved": bool(moved),
            "second": None,
        }
        if has_second:
            m = LEAD_SUFFIX.search(s_project_raw)
            emp["second"] = {
                "sector": s_sector,
                "project": alias.get(s_project, s_project),
                "role": f"{m.group(1)} Lead" if m else None,
                "billing": billing(v(r, "Second Billed Status")),
                "pct": alloc("Second Project %", 0.0) if dual else number(v(r, "Second Project %")),
                "isSeparateAssignment": dual,
            }
        out.append(emp)
    return out


def project_leads(projects, employees):
    """Same rule as the Overall sheet: people marked '<project> — Project Lead' in
    Headcount Data, otherwise the lead on the Projects sheet."""
    for p in projects:
        marked = [e["name"] for e in employees
                  if e["second"] and e["second"]["role"] == "Project Lead"
                  and e["second"]["sector"] == p["sector"] and e["second"]["project"] == p["name"]]
        if marked:
            p["lead"] = ", ".join(marked)


def read_kpi_master(g):
    hdr = [clean(g.cell(1, c)) for c in range(1, 14)]
    rows = []
    for r in range(2, g.max_row + 1):
        if not clean(g.cell(r, 1)):
            continue
        rec = {hdr[c - 1]: clean(g.cell(r, c)) for c in range(1, 14)}
        rows.append({
            "sector": rec["Sector"], "role": rec["Role"], "no": int(number(rec["KPI #"]) or 0),
            "kpi": rec["KPI"], "measure": rec["How measured"], "unit": rec["Unit"],
            "better": rec["Better"], "target": rec["Suggested target"], "weight": rec["Weight"],
            "status": rec["Review status"], "category": rec["Category"], "notes": rec["Definition / target notes"],
        })
    return rows


def read_kpi_matrix(g):
    blocks, cur, area = [], None, None
    for r in range(1, g.max_row + 1):
        vals = [clean(g.cell(r, c)) for c in range(1, 7)]
        a = vals[0]
        if a and (a.startswith("BUSINESS LEAD —") or a.startswith("PROJECT LEAD —")):
            cur = {"title": a, "areas": []}
            blocks.append(cur)
            continue
        if cur is None or a == "Area / KPI" or (a and a.endswith("TOTAL")) or not any(vals):
            continue
        if a:
            area = {"area": a, "weight": vals[1], "metrics": []}
            cur["areas"].append(area)
        if area and vals[2]:
            area["metrics"].append({"metric": vals[2], "weight": vals[3], "target": vals[4], "measure": vals[5]})
    return blocks


def read_pd_kpi(g):
    roles = [(clean(g.cell(1, c)), c) for c in (4, 6, 8, 10, 12)]
    rows, area, totals = [], None, {}
    for r in range(3, g.max_row + 1):
        a, metric = clean(g.cell(r, 1)), clean(g.cell(r, 2))
        if a == "Total weight":
            totals = {role: number(g.cell(r, c + 1)) for role, c in roles}
            break
        if a:
            area = a
        if not metric:
            continue
        rows.append({
            "area": area, "metric": metric, "measure": clean(g.cell(r, 3)),
            "roles": {role: {"target": clean(g.cell(r, c)), "weight": number(g.cell(r, c + 1))} for role, c in roles},
        })
    return {"roles": [r for r, _ in roles], "rows": rows, "totals": totals}


def read_excel_summary(gv, gf):
    """Totals as calculated by the workbook — used only for the accuracy check.
    Cells driven by Google-Sheets-only formulas are flagged, because desktop
    Excel does not recalculate them."""
    levels = [clean(gv.cell(5, c)) for c in range(2, 9)]
    by_sector = {}
    for r in range(6, 14):
        s = clean(gv.cell(r, 1))
        if not s:
            continue
        by_sector[s] = {lvl: gv.cell(r, 2 + i) for i, lvl in enumerate(levels)}
        by_sector[s]["Total"] = gv.cell(r, 9)
    billing_tot = {(clean(gv.cell(r, 1)) or NOT_SET): gv.cell(r, 2) for r in range(18, 23)}
    sector_billing = {}
    for r in range(31, 39):
        s = clean(gv.cell(r, 1))
        if s:
            sector_billing[s] = {
                "Billed": gv.cell(r, 2), "Buffer Billed": gv.cell(r, 3), "Buffer": gv.cell(r, 4),
                "UnBilled": gv.cell(r, 5), NOT_SET: gv.cell(r, 6), "Total": gv.cell(r, 7),
            }
    projects = []
    for r in range(6, 57):
        s, p = clean(gv.cell(r, 11)), clean(gv.cell(r, 12))
        if s and p:
            projects.append({
                "sector": s, "project": p, "assigned": gv.cell(r, 13),
                "Billed": gv.cell(r, 14), "Buffer Billed": gv.cell(r, 15), "Buffer": gv.cell(r, 16),
                "UnBilled": gv.cell(r, 17), NOT_SET: gv.cell(r, 18), "weighted": gv.cell(r, 20),
            })
    google_only = bool(gf) and GOOGLE_ONLY in str(gf.cell(7, 13) or "")
    return {
        "bySectorLevel": by_sector,
        "billing": billing_tot,
        "earning": {"Earning": gv.cell(25, 2), "Not earning": gv.cell(26, 2)},
        "movedSoFar": gv.cell(27, 2),
        "sectorBilling": sector_billing,
        "projects": projects,
        "totalAssignments": gv.cell(57, 13),
        "totalWeighted": gv.cell(57, 20),
        # Project table + totals use Google-Sheets-only formulas (stale if saved by desktop Excel).
        "projectsGoogleOnly": google_only,
    }


# ---------------------------------------------------------------- build
def build(src):
    src = Path(src)
    wb_v = openpyxl.load_workbook(src, data_only=True, read_only=True)
    wb_f = openpyxl.load_workbook(src, data_only=False, read_only=True)
    try:
        hd_v, hd_f = Grid(sheet(wb_v, HEADCOUNT_SHEET)), Grid(sheet(wb_f, HEADCOUNT_SHEET))
        projects_raw = read_projects(Grid(sheet(wb_v, "Projects")))
        employees = read_employees(hd_v, hd_f, projects_raw)
        project_leads(projects_raw, employees)

        sectors = read_operating_model(Grid(sheet(wb_v, "Operating Model")))
        try:
            heads = read_sector_heads(Grid(sheet(wb_v, "Overall")))
        except KeyError:
            heads = {}
        for s in sectors:
            s["head"] = heads.get(s["name"])
        sector_order = {s["name"]: i for i, s in enumerate(sectors)}
        projects = [{
            "sector": p["sector"], "name": p["name"], "lead": p["lead"], "status": p["status"],
            "order": sector_order.get(p["sector"], 50) * 1000 + p["row"],
        } for p in projects_raw]

        def optional(fn, *names):
            try:
                return fn(*[Grid(sheet(wb_v, n)) for n in names])
            except (KeyError, ValueError, TypeError, IndexError):
                return None

        try:
            summary = read_excel_summary(Grid(sheet(wb_v, SUMMARY_SHEET)), Grid(sheet(wb_f, SUMMARY_SHEET)))
        except (KeyError, ValueError, TypeError, IndexError):
            summary = None

        modified = dt.datetime.fromtimestamp(src.stat().st_mtime)
        return {
            "meta": {
                "company": "VConnecTech Systems",
                "source": src.name,
                "sourceModified": modified.isoformat(timespec="seconds"),
                "lastModifiedBy": (wb_v.properties.lastModifiedBy or "").strip() or None,
                "generatedAt": dt.datetime.now().isoformat(timespec="seconds"),
                "mode": "snapshot",
            },
            "sectors": sectors,
            "projects": projects,
            "employees": employees,
            "kpi": {
                "master": optional(read_kpi_master, "KPI_Master") or [],
                "matrix": optional(read_kpi_matrix, "KPI Matrix") or [],
                "productDevelopment": optional(read_pd_kpi, "Product Development KPI"),
            },
            "excelSummary": summary,
        }
    finally:
        wb_v.close()
        wb_f.close()


def write_snapshot(data, out=OUT):
    out.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, indent=1, default=str)
    tmp = out.with_suffix(".tmp")
    tmp.write_text("// Generated by scripts/build_data.py — do not edit by hand.\n"
                   f"window.VCT_DATA = {payload};\n", encoding="utf-8")
    tmp.replace(out)


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else default_workbook()
    data = build(src)
    write_snapshot(data)
    print(f"Wrote {OUT} — {len(data['employees'])} employees, {len(data['projects'])} projects, "
          f"{len(data['kpi']['master'])} KPI rows")


if __name__ == "__main__":
    main()
