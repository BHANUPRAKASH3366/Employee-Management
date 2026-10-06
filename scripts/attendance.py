"""
Attendance alerts from the HRMS export (e.g. "Early_Late Report").

Reads the HRMS workbook set in portal.config.json → "attendance.hrmsPath" and flags employees who
  • had one or more days with fewer than `minHoursPerDay` hours (default 8), and/or
  • attended fewer than `minDaysPerMonth` days in the period (default 20).
`match` = "any" flags people who break either rule; "all" only people who break both.

The flagged list is emailed to HR and the CEO (To) with CC — once per distinct report, so a restart
or re-save with the same result doesn't send it again. A copy of every email is kept in
audit/attendance/outbox. The live server shows the report on the Alerts page (Attendance tab).

Standalone check (prints the report, sends nothing):
    python scripts/attendance.py
"""
import datetime as dt
import hashlib
import html
import json
import os
import re
import shutil
import smtplib
import ssl
import sys
import tempfile
import threading
import time
import traceback
from email.message import EmailMessage
from email.utils import formatdate, make_msgid
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import audit  # noqa: E402
import build_data  # noqa: E402

ATT_DIR = build_data.ROOT / "audit" / "attendance"
REPORT_FILE = ATT_DIR / "report.json"
SENT_FILE = ATT_DIR / "sent.jsonl"
OUTBOX = ATT_DIR / "outbox"


def set_dir(path):
    """Keep attendance records somewhere else (tests)."""
    global ATT_DIR, REPORT_FILE, SENT_FILE, OUTBOX
    ATT_DIR = Path(path)
    REPORT_FILE, SENT_FILE, OUTBOX = ATT_DIR / "report.json", ATT_DIR / "sent.jsonl", ATT_DIR / "outbox"


def load_config():
    cfg = json.loads(build_data.CONFIG.read_text(encoding="utf-8")) if build_data.CONFIG.exists() else {}
    a = dict(cfg.get("attendance", {}))
    a.setdefault("enabled", True)
    a.setdefault("hrmsPath", "")
    a.setdefault("sheet", "")
    a.setdefault("minHoursPerDay", 8)
    a.setdefault("minDaysPerMonth", 20)
    a.setdefault("match", "any")
    a.setdefault("hrEmails", [])
    a.setdefault("ceoEmails", [])
    a.setdefault("ccEmails", [])
    a.setdefault("subjectPrefix", "[VCT Portal] Attendance")
    # Monthly report: emailed once a month on this day/time (sent later that month if the portal was off then).
    mr = dict(a.get("monthlyReport") or {})
    mr.setdefault("enabled", True)
    mr.setdefault("day", 26)
    mr.setdefault("time", "10:00")
    a["monthlyReport"] = mr
    a["smtp"] = audit.load_config().get("smtp", {})  # same sender account as the edit alerts
    a["portalUrl"] = audit.load_config().get("portalUrl", "")
    return a


# ---------------------------------------------------------------- reading the HRMS file
def _hours(v):
    """'9.25' / 9.25 / '08:30' / time → decimal hours. None when blank or '-'."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, dt.time):
        return v.hour + v.minute / 60
    if isinstance(v, dt.timedelta):
        return v.total_seconds() / 3600
    s = str(v).strip()
    m = re.match(r"^(\d+):(\d{1,2})$", s)
    if m:
        return int(m.group(1)) + int(m.group(2)) / 60
    try:
        return float(s)
    except ValueError:
        return None


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v or "").strip()
    for f in ("%d-%b-%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def _text(v):
    if v is None:
        return ""
    if isinstance(v, dt.time):
        return v.strftime("%I:%M %p")
    return str(v).strip()


def read_hrms(path, sheet=""):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        names = wb.sheetnames
        pick = sheet if sheet in names else next((n for n in names if "decimal" in n.lower()), names[0])
        ws = wb[pick]
        head_idx, cols, period = None, {}, {}
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if head_idx is None:
                first = _text(r[0] if r else "").lower()
                if first in ("from date", "to date") and len(r) > 1:
                    period[first] = _date(r[1])
                labels = [_text(c).lower() for c in r]
                if "employee id" in labels and "date" in labels:
                    head_idx = i
                    cols = {k: labels.index(k) for k in labels if k}
                continue
            rows.append(r)
    finally:
        wb.close()
    if head_idx is None:
        raise ValueError(f"No 'Employee Id' / 'Date' header row found in sheet '{pick}'")

    def col(r, name):
        j = cols.get(name)
        return r[j] if j is not None and j < len(r) else None

    emps = {}
    for r in rows:
        eid = _text(col(r, "employee id"))
        d = _date(col(r, "date"))
        if not eid or not d:
            continue
        e = emps.setdefault(eid, {"id": eid, "name": _text(col(r, "employee name")), "email": _text(col(r, "email id")), "days": {}})
        h = _hours(col(r, "total hours"))
        fin, fout = _text(col(r, "first in")), _text(col(r, "last out"))
        fin, fout = ("" if fin == "-" else fin), ("" if fout == "-" else fout)
        e["days"][d.isoformat()] = {"date": d.isoformat(), "in": fin, "out": fout,
                                    "hours": round(h, 2) if h is not None else None,
                                    "missedPunch": bool(fin) != bool(fout) or (not fin and not fout)}
    return pick, period, emps


def build_report(path, cfg):
    sheet, period, emps = read_hrms(path, cfg.get("sheet", ""))
    min_h, min_d = float(cfg["minHoursPerDay"]), int(cfg["minDaysPerMonth"])
    every = all if str(cfg.get("match", "any")).lower() == "all" else any
    people = []
    for e in emps.values():
        days = sorted(e["days"].values(), key=lambda x: x["date"])
        hrs = [d["hours"] or 0 for d in days]
        short = [d for d in days if (d["hours"] or 0) < min_h]
        few = len(days) < min_d
        rules = []
        if short:
            rules.append("short-hours")
        if few:
            rules.append("few-days")
        people.append({
            "id": e["id"], "name": e["name"], "email": e["email"],
            "daysPresent": len(days), "totalHours": round(sum(hrs), 2),
            "avgHours": round(sum(hrs) / len(hrs), 2) if hrs else 0,
            "shortDays": short, "rules": rules,
            "flagged": bool(rules) and every([bool(short), few]),
        })
    people.sort(key=lambda p: (not p["flagged"], p["daysPresent"], -len(p["shortDays"]), p["name"]))
    flagged = [p for p in people if p["flagged"]]
    all_dates = sorted({d for e in emps.values() for d in e["days"]})
    frm = period.get("from date") or (_date(all_dates[0]) if all_dates else None)
    to = period.get("to date") or (_date(all_dates[-1]) if all_dates else None)
    return {
        "source": Path(path).name, "sheet": sheet,
        "periodFrom": frm.isoformat() if frm else None, "periodTo": to.isoformat() if to else None,
        "periodLabel": frm.strftime("%B %Y") if frm and to and (frm.year, frm.month) == (to.year, to.month)
        else (f"{frm:%d %b %Y} – {to:%d %b %Y}" if frm and to else ""),
        "rules": {"minHoursPerDay": min_h, "minDaysPerMonth": min_d, "match": "all" if every is all else "any"},
        "employees": len(people), "workingDatesInFile": len(all_dates),
        "flaggedCount": len(flagged),
        "shortHoursCount": sum(1 for p in people if "short-hours" in p["rules"]),
        "fewDaysCount": sum(1 for p in people if "few-days" in p["rules"]),
        "people": people,
        "builtAt": audit.now_iso(),
    }


def report_key(rep):
    """Same flagged people with the same numbers → same key → not emailed twice."""
    core = [rep["periodFrom"], rep["periodTo"], rep["rules"],
            [(p["id"], p["daysPresent"], [d["date"] for d in p["shortDays"]]) for p in rep["people"] if p["flagged"]]]
    return hashlib.sha1(json.dumps(core, sort_keys=True).encode()).hexdigest()[:16]


# ---------------------------------------------------------------- email
def _fmt_h(h):
    if h is None:
        return "—"
    return f"{int(h)}h {round((h - int(h)) * 60):02d}m"


def build_email(rep, cfg, monthly=False):
    r = rep["rules"]
    flagged = [p for p in rep["people"] if p["flagged"]]
    rule_txt = (f"fewer than {r['minHoursPerDay']:g} hours on a day "
                f"{'and' if r['match'] == 'all' else 'or'} fewer than {r['minDaysPerMonth']} days attended")
    subject = f"{cfg['subjectPrefix']} — {'Monthly report · ' if monthly else ''}{rep['periodLabel'] or 'report'}: {len(flagged)} employee{'s' if len(flagged) != 1 else ''} below the attendance rules"

    lines = [f"Attendance report — {rep['periodLabel']}  ({rep['source']})", "",
             f"Rule: flagged when {rule_txt}.",
             f"Employees in file: {rep['employees']}   Flagged: {len(flagged)}   "
             f"(<{r['minHoursPerDay']:g}h days: {rep['shortHoursCount']}, <{r['minDaysPerMonth']} days: {rep['fewDaysCount']})", ""]
    for p in flagged:
        lines.append(f"{p['id']} — {p['name']}  <{p['email'] or 'no email'}>")
        lines.append(f"    Days attended: {p['daysPresent']}   Total: {_fmt_h(p['totalHours'])}   Average/day: {_fmt_h(p['avgHours'])}")
        if p["shortDays"]:
            lines.append(f"    Days under {r['minHoursPerDay']:g}h ({len(p['shortDays'])}): " +
                         ", ".join(f"{dt.date.fromisoformat(d['date']):%d %b} {_fmt_h(d['hours'])} ({d['in'] or '-'}–{d['out'] or '-'}{', missed punch' if d.get('missedPunch') else ''})" for d in p["shortDays"]))
        lines.append("")
    if cfg.get("portalUrl"):
        lines.append(f"Portal: {cfg['portalUrl'].rstrip('/')}/#/changes?tab=attendance")
    text = "\n".join(lines)

    e = html.escape
    td = 'style="padding:7px 8px;border:1px solid #c9d2de;vertical-align:top"'
    th = 'style="padding:7px 8px;border:1px solid #9fabbd;background:#e6ebf2;text-align:left"'
    rows = []
    for p in flagged:
        few = "few-days" in p["rules"]
        miss = ' <b style="color:#be123c">missed punch</b>'
        short = "; ".join(f"{dt.date.fromisoformat(d['date']):%d %b} — {_fmt_h(d['hours'])} ({e(d['in'] or '-')}–{e(d['out'] or '-')}){miss if d.get('missedPunch') else ''}" for d in p["shortDays"])
        rows.append(f"<tr><td {td}><b>{e(p['id'])}</b></td><td {td}>{e(p['name'])}<br><span style=\"color:#7b879b;font-size:12px\">{e(p['email'])}</span></td>"
                    f"<td {td} align=\"center\"><b style=\"color:{'#be123c' if few else '#15803d'}\">{p['daysPresent']}</b></td>"
                    f"<td {td}>{_fmt_h(p['totalHours'])}</td><td {td}>{_fmt_h(p['avgHours'])}</td>"
                    f"<td {td} align=\"center\"><b style=\"color:{'#be123c' if p['shortDays'] else '#15803d'}\">{len(p['shortDays'])}</b></td>"
                    f"<td {td} style=\"padding:7px 8px;border:1px solid #c9d2de;font-size:12px\">{short or '—'}</td></tr>")
    html_body = f"""<div style="font-family:Segoe UI,Arial,sans-serif;color:#0f1b2d;max-width:980px">
<div style="background:#060b14;color:#fff;padding:14px 18px;border-radius:10px 10px 0 0"><b>VConnecTech Systems</b> · Attendance alert — {e(rep['periodLabel'])}</div>
<div style="border:1px solid #e3e8f0;border-top:0;padding:18px;border-radius:0 0 10px 10px">
<p style="margin:0 0 8px"><b>{len(flagged)}</b> of {rep['employees']} employees are flagged: {e(rule_txt)}.</p>
<p style="margin:0 0 14px;color:#44526a;font-size:13px">Under {r['minHoursPerDay']:g} h on at least one day: <b>{rep['shortHoursCount']}</b> · Fewer than {r['minDaysPerMonth']} days attended: <b>{rep['fewDaysCount']}</b> · Source: {e(rep['source'])}</p>
<table style="border-collapse:collapse;width:100%;font-size:13px">
<tr><th {th}>Emp ID</th><th {th}>Name</th><th {th}>Days attended</th><th {th}>Total hours</th><th {th}>Avg / day</th><th {th}>Days &lt; {r['minHoursPerDay']:g} h</th><th {th}>Short days (date — hours, in–out)</th></tr>
{''.join(rows)}
</table>
{f'<p style="margin-top:16px"><a href="{e(cfg["portalUrl"].rstrip("/"))}/#/changes?tab=attendance">Open attendance alerts in the portal</a></p>' if cfg.get("portalUrl") else ''}
</div></div>"""
    return subject, text, html_body


def send_report(rep, cfg, reason="new report", schedule=None):
    subject, text, html_body = build_email(rep, cfg, monthly=bool(schedule))
    smtp = cfg.get("smtp", {})
    to = [a for a in list(cfg.get("hrEmails", [])) + list(cfg.get("ceoEmails", [])) if a]
    cc = [a for a in cfg.get("ccEmails", []) if a and a not in to]
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = smtp.get("from") or smtp.get("username") or "vct-portal@localhost"
    msg["To"] = ", ".join(to) or "hr-and-ceo (not configured)"
    if cc:
        msg["Cc"] = ", ".join(cc)
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain="vct-portal")
    msg.set_content(text)
    msg.add_alternative(html_body, subtype="html")

    sid = dt.datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    OUTBOX.mkdir(parents=True, exist_ok=True)
    (OUTBOX / f"{sid}.eml").write_bytes(bytes(msg))
    rec = {"id": sid, "at": audit.now_iso(), "key": report_key(rep), "period": rep["periodLabel"], "flagged": rep["flaggedCount"], "schedule": schedule,
           "subject": subject, "to": to, "cc": cc, "reason": reason, "file": f"{sid}.eml"}
    if not cfg.get("enabled", True):
        rec.update(status="off", detail="Attendance alerts are switched off in portal.config.json.")
    elif not to:
        rec.update(status="not-configured", detail="No HR / CEO email set (attendance.hrEmails / ceoEmails). Saved to audit/attendance/outbox.")
    elif not smtp.get("host"):
        rec.update(status="not-configured", detail="No email server set (alerts.smtp). Saved to audit/attendance/outbox.")
    else:
        password = os.environ.get(smtp.get("passwordEnv", "VCT_SMTP_PASSWORD"), "") or smtp.get("password", "")
        security = (smtp.get("security") or "starttls").lower()
        port = int(smtp.get("port") or (465 if security == "ssl" else 587))
        last = None
        for attempt in range(3):
            try:
                ctx = ssl.create_default_context()
                server = smtplib.SMTP_SSL(smtp["host"], port, context=ctx, timeout=30) if security == "ssl" \
                    else smtplib.SMTP(smtp["host"], port, timeout=30)
                with server:
                    if security == "starttls":
                        server.starttls(context=ctx)
                    if smtp.get("username"):
                        server.login(smtp["username"], password)
                    server.send_message(msg, to_addrs=to + cc)
                last = None
                break
            except Exception as ex:
                last = ex
                time.sleep(2 * (attempt + 1))
        rec.update(status="failed", detail=f"Email failed: {last}") if last else \
            rec.update(status="sent", detail="Sent to " + ", ".join(to) + (" · CC " + ", ".join(cc) if cc else ""))
    with SENT_FILE.open("a", encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    return rec


EMAIL_RE = re.compile(r"^[^@\s,;<>]+@[^@\s,;<>]+\.[^@\s,;<>]+$")


def save_recipients(hr, ceo, cc):
    """Update attendance.hrEmails / ceoEmails / ccEmails in portal.config.json (from the portal page).
    Returns (ok, message). Other settings in the file are left as they are."""
    lists = {}
    for key, vals in (("hrEmails", hr), ("ceoEmails", ceo), ("ccEmails", cc)):
        if not isinstance(vals, list):
            return False, f"{key} must be a list of email addresses."
        clean = []
        for v in vals:
            v = str(v or "").strip()
            if not v:
                continue
            if not EMAIL_RE.match(v) or len(v) > 254:
                return False, f"“{v}” is not a valid email address."
            if v.lower() not in [c.lower() for c in clean]:
                clean.append(v)
        if len(clean) > 50:
            return False, "Too many addresses (50 per list at most)."
        lists[key] = clean
    if not lists["hrEmails"] and not lists["ceoEmails"]:
        return False, "Add at least one HR or CEO address."
    cfg_path = build_data.CONFIG
    cfg = json.loads(cfg_path.read_text(encoding="utf-8")) if cfg_path.exists() else {}
    cfg.setdefault("attendance", {}).update(lists)
    tmp = cfg_path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(cfg, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    os.replace(tmp, cfg_path)  # all-or-nothing, so a half-written config can never be read
    return True, "Recipients saved."


def read_sent(limit=50):
    if not SENT_FILE.exists():
        return []
    out = []
    for line in SENT_FILE.read_text(encoding="utf-8").splitlines():
        try:
            out.append(json.loads(line))
        except ValueError:
            pass
    return out[::-1][:limit]


# ---------------------------------------------------------------- monthly schedule
def _monthly_due(mr, year, month):
    """The send moment for a month: configured day (capped at the month's last day) and time."""
    nxt = dt.date(year + (month == 12), month % 12 + 1, 1)
    last = (nxt - dt.timedelta(days=1)).day
    try:
        hh, mm = (int(x) for x in str(mr.get("time", "10:00")).split(":")[:2])
    except ValueError:
        hh, mm = 10, 0
    return dt.datetime(year, month, min(max(1, int(mr.get("day", 26))), last), hh, mm)


def monthly_info(cfg, now=None):
    """{"enabled", "day", "time", "next", "lastSent"} for the portal page."""
    mr = cfg["monthlyReport"]
    now = now or dt.datetime.now()
    tag = now.strftime("%Y-%m")
    sent = [s for s in read_sent(500) if s.get("schedule")]
    done = any(s["schedule"] == tag and s.get("status") in ("sent", "not-configured", "off") for s in sent)
    due = _monthly_due(mr, now.year, now.month)
    if done:
        due = _monthly_due(mr, now.year + (now.month == 12), now.month % 12 + 1)
    last = next((s for s in sent if s.get("status") == "sent"), None)
    return {"enabled": bool(mr.get("enabled")), "day": mr.get("day"), "time": mr.get("time"),
            "next": due.isoformat(timespec="minutes") if mr.get("enabled") else None,
            "dueNow": bool(mr.get("enabled")) and not done and now >= _monthly_due(mr, now.year, now.month),
            "lastSent": last["at"] if last else None}


# ---------------------------------------------------------------- watcher (used by live_server.py)
class AttendanceWatcher:
    """Watches the HRMS file; rebuilds the report when it changes and emails new results once."""

    def __init__(self, log=print, interval=5):
        self.log = log
        self.interval = interval
        self.lock = threading.Lock()
        self.report = None
        self.error = None
        self._seen = None
        self._pending = None
        if REPORT_FILE.exists():
            try:
                self.report = json.loads(REPORT_FILE.read_text(encoding="utf-8"))
            except ValueError:
                pass

    def path(self):
        p = load_config().get("hrmsPath") or ""
        return Path(p) if p else None

    def check(self):
        p = self.path()
        cfg_sig = self._config_sig()
        if cfg_sig != getattr(self, "_cfg_sig", cfg_sig):
            self._seen = None  # rules or recipients changed in portal.config.json → re-check (and email if new)
        self._cfg_sig = cfg_sig
        if not p:
            self.error = "No HRMS file set (attendance.hrmsPath in portal.config.json)."
            return
        try:
            st = p.stat()
        except FileNotFoundError:
            self.error = f"HRMS file not found: {p}"
            return
        sig = (str(p), st.st_mtime_ns, st.st_size)
        if sig == self._seen:
            return
        if self._pending != sig:  # wait one round so a file still being saved isn't read
            self._pending = sig
            return
        self._seen, self._pending = sig, None
        try:
            self.refresh(p)
        except PermissionError:
            self._seen = None
        except Exception as ex:
            self.error = f"Could not read the HRMS file: {ex}"
            self.log("Attendance: " + self.error + "\n" + traceback.format_exc())

    def refresh(self, p, send=True):
        cfg = load_config()
        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp) / p.name
            shutil.copy2(p, copy)
            rep = build_report(copy, cfg)
        rep["source"] = p.name
        rep["sourceModified"] = dt.datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds")
        ATT_DIR.mkdir(parents=True, exist_ok=True)
        REPORT_FILE.write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
        with self.lock:
            self.report, self.error = rep, None
        self.log(f"Attendance: {rep['periodLabel']} — {rep['flaggedCount']} of {rep['employees']} employees flagged")
        # Email each distinct result once per set of recipients (adding HR/CEO later still sends it).
        to = [a for a in list(cfg["hrEmails"]) + list(cfg["ceoEmails"]) if a]
        cc = [a for a in cfg["ccEmails"] if a and a not in to]
        done = any(s.get("key") == report_key(rep) and s.get("to") == to and s.get("cc") == cc
                   and s.get("status") in ("sent", "not-configured", "off") for s in read_sent(500))
        if send and rep["flaggedCount"] and not done:
            res = send_report(rep, cfg)
            self.log(f"Attendance email: {res['status']} — {res['detail']}")

    @staticmethod
    def _config_sig():
        return json.dumps([load_config().get(k) for k in ("hrEmails", "ceoEmails", "ccEmails", "minHoursPerDay", "minDaysPerMonth", "match", "sheet")])

    def update_recipients(self, hr, ceo, cc):
        """Save from the portal. Changing who receives emails doesn't by itself resend the current report —
        use "Send report again" for that; the monthly report and new HRMS files go to the new list."""
        ok, msg = save_recipients(hr, ceo, cc)
        if ok:
            self._cfg_sig = self._config_sig()
            self.log("Attendance recipients updated from the portal")
        return ok, msg

    def send_now(self):
        with self.lock:
            rep = self.report
        if not rep:
            return {"status": "failed", "detail": "No attendance report yet."}
        return send_report(rep, load_config(), reason="sent again from the portal")

    def check_monthly(self):
        """On/after the configured day each month, email the report once (retry failures every 30 min)."""
        cfg = load_config()
        info = monthly_info(cfg)
        if not cfg.get("enabled", True) or not info["dueNow"]:
            return
        if time.time() - getattr(self, "_monthly_try", 0) < 1800:
            return
        self._monthly_try = time.time()
        p = self.path()
        if p and p.exists():
            try:
                self.refresh(p, send=False)  # use the latest HRMS file
            except Exception as ex:
                self.log(f"Attendance monthly report: could not re-read the HRMS file ({ex}); using the last report")
        with self.lock:
            rep = self.report
        if not rep:
            self.log("Attendance monthly report: no report to send yet (HRMS file not read)")
            return
        tag = dt.datetime.now().strftime("%Y-%m")
        res = send_report(rep, cfg, reason=f"monthly report ({dt.datetime.now():%B %Y})", schedule=tag)
        self.log(f"Attendance monthly report: {res['status']} — {res['detail']}")

    def watch(self):
        while True:
            try:
                self.check()
                self.check_monthly()
            except Exception:
                self.log("Attendance watcher error:\n" + traceback.format_exc())
            time.sleep(self.interval)

    def status(self):
        cfg = load_config()
        smtp = cfg.get("smtp", {})
        to = [a for a in list(cfg["hrEmails"]) + list(cfg["ceoEmails"]) if a]
        with self.lock:
            return {
                "report": self.report, "error": self.error, "sent": read_sent(),
                "config": {"enabled": cfg["enabled"], "hrmsPath": cfg["hrmsPath"], "hrEmails": cfg["hrEmails"],
                           "ceoEmails": cfg["ceoEmails"], "ccEmails": cfg["ccEmails"],
                           "emailConfigured": bool(to) and bool(smtp.get("host"))},
                "monthly": monthly_info(cfg),
            }


if __name__ == "__main__":
    c = load_config()
    path = sys.argv[1] if len(sys.argv) > 1 else c["hrmsPath"]
    rep = build_report(path, c)
    print(f"{rep['periodLabel']}: {rep['flaggedCount']} of {rep['employees']} flagged "
          f"(<{rep['rules']['minHoursPerDay']:g}h days: {rep['shortHoursCount']}, <{rep['rules']['minDaysPerMonth']} days: {rep['fewDaysCount']})")
    for p in rep["people"]:
        if p["flagged"]:
            print(f"  {p['id']:<10} {p['name']:<34} days {p['daysPresent']:>2}  short {len(p['shortDays']):>2}  avg {p['avgHours']}")
